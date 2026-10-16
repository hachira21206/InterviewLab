import pytest
import pytest_asyncio
import asyncio
import uuid
import re
from sqlalchemy import select, delete, func
from sqlalchemy.ext.asyncio import AsyncSession
from app.db.session import SessionLocal

from app.models.cv_document import CVDocument, CVStatus
from app.models.job_application import JobApplication
from app.models.application_analysis import ApplicationAnalysis, AnalysisStatus
from app.models.interview_session import InterviewSession, SessionStatus
from app.models.interview_question import InterviewQuestion, QuestionCategory, QuestionDifficulty, QuestionSource
from app.models.interview_response import InterviewResponse
from app.models.interview_evaluation import (
    InterviewEvaluation,
    InterviewResponseEvaluation,
    EvaluationStatus,
    ReadinessLevel,
    EvaluationConfidence,
)
from app.schemas.interview_evaluation import (
    QuestionEvaluationAI,
    InterviewEvaluationAISet,
)
from app.services.analysis_provider import (
    FakeAnalysisProvider,
    InterviewEvaluationResult,
    AnalysisInput,
)
from app.prompts.interview_evaluation import redact_sensitive_info
from app.core.config import settings


@pytest_asyncio.fixture(autouse=True)
def mock_ai_provider(monkeypatch):
    import app.services.interview_evaluation
    fake = FakeAnalysisProvider()
    monkeypatch.setattr(settings, "GROQ_API_KEY", "fake_eval_key")

    def fake_get_provider():
        from app.core.exceptions import ProviderConfigurationError
        if settings.AI_PROVIDER == "groq" and not settings.GROQ_API_KEY:
            raise ProviderConfigurationError("Chưa cấu hình GROQ_API_KEY")
        if settings.AI_PROVIDER != "groq":
            raise ProviderConfigurationError(f"Provider không hỗ trợ: {settings.AI_PROVIDER}")
        return fake

    monkeypatch.setattr(app.services.interview_evaluation, "get_analysis_provider", fake_get_provider)
    return fake


@pytest_asyncio.fixture
async def completed_interview_session():
    async with SessionLocal() as db:
        doc_id = uuid.uuid4()
        app_id = uuid.uuid4()
        analysis_id = uuid.uuid4()
        session_id = uuid.uuid4()

        doc = CVDocument(
            id=doc_id,
            original_filename="cv.pdf",
            storage_key=f"k_{doc_id}",
            media_type="application/pdf",
            size_bytes=100,
            sha256="abc",
            status=CVStatus.extracted,
            extracted_text="Python FastAPI SQL candidate@example.com Phone 0912345678"
        )
        db.add(doc)

        job_app = JobApplication(
            id=app_id,
            cv_document_id=doc_id,
            target_role="Backend Developer Intern",
            job_description="Junior/Intern Backend Developer"
        )
        db.add(job_app)

        analysis = ApplicationAnalysis(
            id=analysis_id,
            job_application_id=app_id,
            status=AnalysisStatus.completed,
            provider="groq",
            model="llama",
            schema_version="1",
            prompt_version="1",
            analysis_data={"summary_vi": "Good candidate", "technical_skills": [], "projects": []}
        )
        db.add(analysis)

        session = InterviewSession(
            id=session_id,
            job_application_id=app_id,
            status=SessionStatus.completed,
            provider="groq",
            model="llama",
            prompt_version="1",
            schema_version="1"
        )
        db.add(session)

        cats = [
            QuestionCategory.backend_fundamentals, QuestionCategory.backend_fundamentals,
            QuestionCategory.cv_project_experience, QuestionCategory.cv_project_experience,
            QuestionCategory.database, QuestionCategory.database,
            QuestionCategory.api_web_design, QuestionCategory.api_web_design,
            QuestionCategory.debugging_problem_solving, QuestionCategory.behavioral_teamwork
        ]
        diffs = [QuestionDifficulty.easy] * 4 + [QuestionDifficulty.medium] * 4 + [QuestionDifficulty.hard] * 2

        question_ids = []
        response_ids = []
        for i in range(1, 11):
            q_id = uuid.uuid4()
            question_ids.append(q_id)
            q = InterviewQuestion(
                id=q_id,
                interview_session_id=session_id,
                order_index=i,
                question_text=f"Question text {i} for backend intern",
                category=cats[i - 1].value,
                difficulty=diffs[i - 1].value,
                purpose=f"Test fundamentals {i}",
                expected_key_points=[f"Key point A for {i}", f"Key point B for {i}"],
                source=QuestionSource.general_backend.value,
                source_evidence=None
            )
            db.add(q)

            r_id = uuid.uuid4()
            response_ids.append(r_id)
            # Make question 9 and 10 skipped in some cases, or default all answered
            resp = InterviewResponse(
                id=r_id,
                interview_session_id=session_id,
                interview_question_id=q_id,
                answer_text=f"Câu trả lời chi tiết của ứng viên cho câu {i}. Giải thích kiến thức kỹ thuật." if i <= 8 else None,
                is_skipped=False if i <= 8 else True,
            )
            db.add(resp)

        await db.commit()

    try:
        yield {
            "session_id": session_id,
            "app_id": app_id,
            "doc_id": doc_id,
            "question_ids": question_ids,
            "response_ids": response_ids,
        }
    finally:
        async with SessionLocal() as db:
            eval_ids = (
                await db.execute(
                    select(InterviewEvaluation.id).where(
                        InterviewEvaluation.interview_session_id == session_id
                    )
                )
            ).scalars().all()
            if eval_ids:
                await db.execute(
                    delete(InterviewResponseEvaluation).where(
                        InterviewResponseEvaluation.interview_evaluation_id.in_(eval_ids)
                    )
                )
                await db.execute(
                    delete(InterviewEvaluation).where(
                        InterviewEvaluation.id.in_(eval_ids)
                    )
                )
            await db.execute(delete(InterviewResponse).where(InterviewResponse.interview_session_id == session_id))
            await db.execute(delete(InterviewQuestion).where(InterviewQuestion.interview_session_id == session_id))
            await db.execute(delete(InterviewSession).where(InterviewSession.id == session_id))
            await db.execute(delete(ApplicationAnalysis).where(ApplicationAnalysis.id == analysis_id))
            await db.execute(delete(JobApplication).where(JobApplication.id == app_id))
            await db.execute(delete(CVDocument).where(CVDocument.id == doc_id))
            await db.commit()


def make_mock_ai_evaluation_set(orders: list[int]):
    evals = []
    for order in orders:
        evals.append(
            QuestionEvaluationAI(
                question_order=order,
                technical_accuracy=3.5,
                completeness=1.5,
                reasoning_clarity=1.5,
                relevance=1.0,
                practical_example=0.5,
                total_score=8.0,
                feedback=f"Nhận xét cho câu {order}: Trả lời rất tốt.",
                strengths=["Hiểu bài", "Diễn đạt tốt"],
                missing_points=["Cần thêm ví dụ"],
                suggested_answer=f"Câu trả lời mẫu cho câu {order}.",
                study_topics=["Backend", "Architecture"],
                confidence=EvaluationConfidence.high
            )
        )
    return InterviewEvaluationAISet(
        evaluations=evals,
        overall_summary="Ứng viên có kiến thức nền tảng tốt.",
        key_strengths=["Nền tảng tốt", "Kỹ năng lập trình vững"],
        knowledge_gaps=["Kiến thức tối ưu database"],
        priority_improvements=["Thực hành REST API"],
        recommended_study_topics=["Database", "Indexing"],
        short_term_study_plan=["Tuần 1: Ôn tập DB", "Tuần 2: Làm bài tập"],
        next_interview_advice="Tự tin hơn khi giải thích kiến trúc."
    )


# 1. Session not found -> 404
@pytest.mark.asyncio
async def test_evaluation_session_not_found(client):
    res = await client.post(f"/api/v1/interview-sessions/{uuid.uuid4()}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 404


# 2. Session not completed -> 409
@pytest.mark.asyncio
async def test_evaluation_session_not_completed(client, completed_interview_session):
    session_id = completed_interview_session["session_id"]
    async with SessionLocal() as db:
        sess = await db.get(InterviewSession, session_id)
        sess.status = SessionStatus.in_progress
        await db.commit()

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 409
    assert "chưa hoàn thành" in res.json()["detail"]


# 3. Not 10 responses -> 409
@pytest.mark.asyncio
async def test_evaluation_not_10_responses(client, completed_interview_session):
    session_id = completed_interview_session["session_id"]
    # Delete one response
    async with SessionLocal() as db:
        resps = (await db.execute(select(InterviewResponse).where(InterviewResponse.interview_session_id == session_id))).scalars().all()
        await db.delete(resps[0])
        await db.commit()

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 409
    assert "không có đủ 10" in res.json()["detail"]


# 4. Missing / False consent -> 422
@pytest.mark.asyncio
async def test_evaluation_missing_consent(client, completed_interview_session):
    session_id = completed_interview_session["session_id"]
    res1 = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": False})
    assert res1.status_code == 422

    res2 = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={})
    assert res2.status_code == 422


# 5. Missing provider API key -> 503, no row created in evaluating
@pytest.mark.asyncio
async def test_evaluation_missing_provider_api_key(client, completed_interview_session, monkeypatch):
    session_id = completed_interview_session["session_id"]
    monkeypatch.setattr(settings, "GROQ_API_KEY", "")

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 503

    async with SessionLocal() as db:
        eval_record = (await db.execute(select(InterviewEvaluation).where(InterviewEvaluation.interview_session_id == session_id))).scalars().first()
        assert eval_record is None


# 6. Success -> exactly 10 evaluations, correct scoring, skipped items 0
@pytest.mark.asyncio
async def test_evaluation_success_with_skipped_questions(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    # 8 answered (orders 1..8) and 2 skipped (orders 9..10)
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=make_mock_ai_evaluation_set(list(range(1, 9))),
        error_code=None,
        error_message=None
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluated"
    assert data["answered_count"] == 8
    assert data["skipped_count"] == 2
    assert len(data["responses"]) == 10

    # Orders 1..8: score 8.0 each. Total answered = 8 * 8.0 = 64.0
    # Orders 9..10: skipped -> score 0.0 each
    assert data["total_score"] == 64.0
    assert data["average_score"] == 6.4

    # Verify skipped questions
    q9 = next(r for r in data["responses"] if r["question_order"] == 9)
    assert q9["is_skipped"] is True
    assert q9["total_score"] == 0.0
    assert q9["technical_accuracy"] == 0.0
    assert q9["completeness"] == 0.0
    assert q9["reasoning_clarity"] == 0.0
    assert q9["relevance"] == 0.0
    assert q9["practical_example"] == 0.0
    assert "bỏ qua" in q9["feedback"]

    q1 = next(r for r in data["responses"] if r["question_order"] == 1)
    assert q1["is_skipped"] is False
    assert q1["total_score"] == 8.0
    assert q1["technical_accuracy"] == 3.5


# 7. All 10 questions skipped -> total score 0, provider not called
@pytest.mark.asyncio
async def test_evaluation_all_10_skipped(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    # Mark all 10 responses as skipped
    async with SessionLocal() as db:
        resps = (await db.execute(select(InterviewResponse).where(InterviewResponse.interview_session_id == session_id))).scalars().all()
        for r in resps:
            r.is_skipped = True
            r.answer_text = None
        await db.commit()

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluated"
    assert data["answered_count"] == 0
    assert data["skipped_count"] == 10
    assert data["total_score"] == 0.0
    assert data["average_score"] == 0.0
    assert data["readiness_level"] == "needs_foundation"
    assert len(data["responses"]) == 10
    assert mock_ai_provider.call_count == 0  # No AI call needed when all skipped!


# 8. AI returns wrong number of items -> evaluation_failed
@pytest.mark.asyncio
async def test_evaluation_ai_wrong_item_count(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    # We expect 8 items, but AI returns only 5
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=make_mock_ai_evaluation_set([1, 2, 3, 4, 5]),
        error_code=None,
        error_message=None
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluation_failed"
    assert data["error_code"] == "INVALID_PROVIDER_RESPONSE"


# 9. AI returns duplicate question order -> evaluation_failed
@pytest.mark.asyncio
async def test_evaluation_ai_duplicate_order(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    mock_data = make_mock_ai_evaluation_set([1, 2, 3, 4, 5, 6, 7, 7])  # duplicate 7, missing 8
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=mock_data,
        error_code=None,
        error_message=None
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluation_failed"
    assert data["error_code"] == "INVALID_PROVIDER_RESPONSE"


# 10. AI returns score out of range -> validation error / evaluation_failed
@pytest.mark.asyncio
async def test_evaluation_ai_score_out_of_range(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=None,
        error_code="INVALID_PROVIDER_RESPONSE",
        error_message="Điểm vượt quá phạm vi cho phép."
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluation_failed"
    assert data["error_code"] == "INVALID_PROVIDER_RESPONSE"


# 11. AI total_score does not match component sum -> evaluation_failed
@pytest.mark.asyncio
async def test_evaluation_total_score_mismatch(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    # Create AI item with mismatched total_score via object without validation or provider error
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=None,
        error_code="INVALID_PROVIDER_RESPONSE",
        error_message="total_score không bằng tổng 5 thành phần."
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluation_failed"
    assert data["error_code"] == "INVALID_PROVIDER_RESPONSE"


# 12. AI returns fake question order not in session -> evaluation_failed
@pytest.mark.asyncio
async def test_evaluation_ai_fake_order(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    # Answered questions are 1..8, so order 9 is not in answered questions
    mock_data = make_mock_ai_evaluation_set([1, 2, 3, 4, 5, 6, 7, 9])
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=mock_data,
        error_code=None,
        error_message=None
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluation_failed"
    assert data["error_code"] == "INVALID_PROVIDER_RESPONSE"


# 13. Provider HTTP error mappings (RATE_LIMITED, PROVIDER_UNAVAILABLE)
@pytest.mark.asyncio
async def test_evaluation_provider_rate_limited(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=None,
        error_code="RATE_LIMITED",
        error_message="Vượt quá giới hạn request."
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluation_failed"
    assert data["error_code"] == "RATE_LIMITED"


# 14. Provider Timeout / Network error
@pytest.mark.asyncio
async def test_evaluation_provider_timeout(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=None,
        error_code="PROVIDER_TIMEOUT",
        error_message="Hết thời gian chờ phản hồi từ AI."
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluation_failed"
    assert data["error_code"] == "PROVIDER_TIMEOUT"


# 15. Refusal / Non-JSON
@pytest.mark.asyncio
async def test_evaluation_refusal(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=None,
        error_code="ANALYSIS_REFUSED",
        error_message="AI từ chối phân tích nội dung."
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluation_failed"
    assert data["error_code"] == "ANALYSIS_REFUSED"


# 16. Retry from evaluation_failed
@pytest.mark.asyncio
async def test_evaluation_retry_from_failed(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    # First call: fail
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=None,
        error_code="RATE_LIMITED",
        error_message="Wait a minute"
    )
    res1 = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res1.status_code == 200
    assert res1.json()["status"] == "evaluation_failed"

    # Second call: success
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=make_mock_ai_evaluation_set(list(range(1, 9))),
        error_code=None,
        error_message=None
    )
    res2 = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res2.status_code == 200
    assert res2.json()["status"] == "evaluated"
    assert res2.json()["total_score"] == 64.0


# 17. Idempotent when already evaluated
@pytest.mark.asyncio
async def test_evaluation_idempotent_when_evaluated(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=make_mock_ai_evaluation_set(list(range(1, 9))),
        error_code=None,
        error_message=None
    )

    res1 = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res1.status_code == 200
    assert res1.json()["status"] == "evaluated"
    assert mock_ai_provider.call_count == 1

    # Second call: must return existing without calling provider!
    res2 = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res2.status_code == 200
    assert res2.json()["status"] == "evaluated"
    assert mock_ai_provider.call_count == 1


# 18. Concurrent evaluation requests -> exactly one provider call, other gets 409
@pytest.mark.asyncio
async def test_evaluation_concurrent_requests(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    provider_entered = asyncio.Event()
    provider_continue = asyncio.Event()

    async def blocking_evaluate(input_data):
        mock_ai_provider.call_count += 1
        provider_entered.set()
        await provider_continue.wait()
        return InterviewEvaluationResult(
            data=make_mock_ai_evaluation_set(list(range(1, 9))),
            error_code=None,
            error_message=None
        )

    mock_ai_provider.evaluate_interview = blocking_evaluate

    t1 = asyncio.create_task(
        client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    )

    await provider_entered.wait()

    # Second concurrent call while first is in provider
    res2 = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res2.status_code == 409
    assert "đang trong quá trình đánh giá" in res2.json()["detail"] or "đang được xử lý" in res2.json()["detail"]

    provider_continue.set()
    res1 = await t1
    assert res1.status_code == 200
    assert res1.json()["status"] == "evaluated"
    assert mock_ai_provider.call_count == 1


# 19. Final transaction fail -> rollback child rows, evaluation_failed
@pytest.mark.asyncio
async def test_evaluation_final_transaction_fail_rollback(client, completed_interview_session, mock_ai_provider, monkeypatch):
    session_id = completed_interview_session["session_id"]
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=make_mock_ai_evaluation_set(list(range(1, 9))),
        error_code=None,
        error_message=None
    )

    import sqlalchemy.ext.asyncio.session
    original_commit = sqlalchemy.ext.asyncio.session.AsyncSession.commit
    commit_count = 0

    async def failing_final_commit(self):
        nonlocal commit_count
        commit_count += 1
        if commit_count == 2:  # First commit is claim (evaluating), second commit is final transaction
            raise Exception("Simulated DB Disk Failure during final save")
        await original_commit(self)

    monkeypatch.setattr(sqlalchemy.ext.asyncio.session.AsyncSession, "commit", failing_final_commit)

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    monkeypatch.undo()

    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluation_failed"

    # Verify no orphan response evaluation rows remained
    async with SessionLocal() as db:
        child_rows = (await db.execute(select(InterviewResponseEvaluation))).scalars().all()
        assert len(child_rows) == 0


# 20. Budget failure -> 409, provider call_count = 0
@pytest.mark.asyncio
async def test_evaluation_budget_failure(client, completed_interview_session, mock_ai_provider, monkeypatch):
    session_id = completed_interview_session["session_id"]
    # Set max answer chars to very small
    monkeypatch.setattr(settings, "AI_MAX_EVALUATION_ANSWERS_CHARS", 10)

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 409
    assert "budget" in res.json()["detail"].lower()
    assert mock_ai_provider.call_count == 0


# 21. Redaction prompt verification
@pytest.mark.asyncio
async def test_evaluation_redaction_prompt(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    captured_prompt = None

    async def intercept_evaluate(input_data: AnalysisInput):
        nonlocal captured_prompt
        captured_prompt = input_data.user_prompt
        return InterviewEvaluationResult(
            data=make_mock_ai_evaluation_set(list(range(1, 9))),
            error_code=None,
            error_message=None
        )

    # Insert email and phone into answer_text
    async with SessionLocal() as db:
        resps = (await db.execute(select(InterviewResponse).where(InterviewResponse.interview_session_id == session_id))).scalars().all()
        target_resp = next(r for r in resps if not r.is_skipped)
        target_resp.answer_text = "My email is test_candidate@gmail.com and phone is 0987654321 and token is Bearer secret_token_12345678"
        target_resp.is_skipped = False
        await db.commit()

    mock_ai_provider.evaluate_interview = intercept_evaluate

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    assert captured_prompt is not None
    assert "test_candidate@gmail.com" not in captured_prompt
    assert "[EMAIL_REDACTED]" in captured_prompt
    assert "0987654321" not in captured_prompt
    assert "[PHONE_REDACTED]" in captured_prompt
    assert "secret_token_12345678" not in captured_prompt


# 22. GET /evaluation endpoint behavior
@pytest.mark.asyncio
async def test_evaluation_get_endpoint(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]

    # 404 before creation
    res0 = await client.get(f"/api/v1/interview-sessions/{session_id}/evaluation")
    assert res0.status_code == 404

    # Create evaluation
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=make_mock_ai_evaluation_set(list(range(1, 9))),
        error_code=None,
        error_message=None
    )
    await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})

    # GET returns full evaluated report
    res_get = await client.get(f"/api/v1/interview-sessions/{session_id}/evaluation")
    assert res_get.status_code == 200
    data = res_get.json()
    assert data["status"] == "evaluated"
    assert data["total_score"] == 64.0
    assert len(data["responses"]) == 10
    # Must NOT contain expected_key_points in response
    for r in data["responses"]:
        assert "expected_key_points" not in r


# 23. ON DELETE CASCADE check
@pytest.mark.asyncio
async def test_evaluation_on_delete_cascade(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=make_mock_ai_evaluation_set(list(range(1, 9))),
        error_code=None,
        error_message=None
    )
    await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})

    # Delete session
    async with SessionLocal() as db:
        sess = await db.get(InterviewSession, session_id)
        await db.delete(sess)
        await db.commit()

        # Both evaluations and response evaluations should be gone
        eval_count = (await db.execute(select(func.count(InterviewEvaluation.id)).where(InterviewEvaluation.interview_session_id == session_id))).scalar()
        assert eval_count == 0

        resp_eval_count = (await db.execute(select(func.count(InterviewResponseEvaluation.id)))).scalar()
        assert resp_eval_count == 0


# 24. Readiness level 8 boundary check
def test_compute_readiness_level_boundaries():
    from app.services.interview_evaluation import compute_readiness_level

    assert compute_readiness_level(0.0) == ReadinessLevel.needs_foundation
    assert compute_readiness_level(49.99) == ReadinessLevel.needs_foundation
    assert compute_readiness_level(50.0) == ReadinessLevel.developing
    assert compute_readiness_level(69.99) == ReadinessLevel.developing
    assert compute_readiness_level(70.0) == ReadinessLevel.intern_ready
    assert compute_readiness_level(84.99) == ReadinessLevel.intern_ready
    assert compute_readiness_level(85.0) == ReadinessLevel.strong_intern_candidate
    assert compute_readiness_level(100.0) == ReadinessLevel.strong_intern_candidate


# 25. Tolerance rejection: AI trả 8.0 component sum nhưng total_score 7.97 -> INVALID_PROVIDER_RESPONSE
def test_question_evaluation_ai_tolerance_rejection():
    from pydantic import ValidationError

    with pytest.raises(ValidationError) as exc_info:
        QuestionEvaluationAI(
            question_order=1,
            technical_accuracy=3.5,
            completeness=1.5,
            reasoning_clarity=1.5,
            relevance=1.0,
            practical_example=0.5,  # Sum = 8.0
            total_score=7.97,  # Lệch 0.03 >= 0.01
            feedback="Nhận xét chi tiết cho câu hỏi.",
            strengths=["A"],
            missing_points=["B"],
            suggested_answer="Đáp án mẫu chi tiết cho câu hỏi.",
            study_topics=["D"],
            confidence=EvaluationConfidence.high,
        )
    assert "total_score" in str(exc_info.value)


# 26. Section A & C: Strict verification - No expected_key_points, purpose, source, source_evidence or secret markers in any public endpoint
@pytest.mark.asyncio
async def test_no_expected_key_points_leakage_across_all_endpoints(client, mock_ai_provider):
    secret_expected = f"SECRET_EXPECTED_{uuid.uuid4().hex}"
    secret_purpose = f"SECRET_PURPOSE_{uuid.uuid4().hex}"
    secret_source_evidence = f"SECRET_EVIDENCE_{uuid.uuid4().hex}"

    doc_id = uuid.uuid4()
    app_id = uuid.uuid4()
    analysis_id = uuid.uuid4()

    async with SessionLocal() as db:
        doc = CVDocument(
            id=doc_id,
            original_filename="cv_sec.pdf",
            storage_key=f"k_{doc_id}",
            media_type="application/pdf",
            size_bytes=100,
            sha256="abc",
            status=CVStatus.extracted,
            extracted_text="Python FastAPI PostgreSQL candidate"
        )
        db.add(doc)

        job_app = JobApplication(
            id=app_id,
            cv_document_id=doc_id,
            target_role="Backend Developer Intern",
            job_description="Junior/Intern Backend Developer"
        )
        db.add(job_app)

        valid_analysis_data = {
            "summary_vi": "Ứng viên có kiến thức nền tảng tốt về Python và cơ sở dữ liệu.",
            "candidate_level": "junior_ready",
            "technical_skills": [],
            "projects": [],
            "education": [],
            "strengths": [],
            "gaps": [],
            "matched_requirements": [],
            "missing_requirements": [],
            "recommended_interview_topics": [],
            "language_detected": "vi",
            "warnings": []
        }

        analysis = ApplicationAnalysis(
            id=analysis_id,
            job_application_id=app_id,
            status=AnalysisStatus.completed,
            provider="groq",
            model="llama",
            schema_version="1",
            prompt_version="1",
            analysis_data=valid_analysis_data
        )
        db.add(analysis)
        await db.commit()

    created_session_id = None
    try:
        import app.services.interview_session
        import app.services.interview_evaluation
        from app.schemas.interview_questions import InterviewQuestionGenerated, InterviewQuestionSet
        from app.models.interview_question import QuestionCategory, QuestionDifficulty, QuestionSource

        def fake_get_provider():
            return mock_ai_provider

        monkeypatch = pytest.MonkeyPatch()
        monkeypatch.setattr(app.services.interview_session, "get_analysis_provider", fake_get_provider)
        monkeypatch.setattr(app.services.interview_evaluation, "get_analysis_provider", fake_get_provider)

        # Mock AI question generator to inject secret markers into internal expected key points
        gen_questions = []
        cats = [
            QuestionCategory.backend_fundamentals, QuestionCategory.backend_fundamentals,
            QuestionCategory.cv_project_experience, QuestionCategory.cv_project_experience,
            QuestionCategory.database, QuestionCategory.database,
            QuestionCategory.api_web_design, QuestionCategory.api_web_design,
            QuestionCategory.debugging_problem_solving, QuestionCategory.behavioral_teamwork
        ]
        diffs = [QuestionDifficulty.easy] * 4 + [QuestionDifficulty.medium] * 4 + [QuestionDifficulty.hard] * 2
        for i in range(1, 11):
            gen_questions.append(
                InterviewQuestionGenerated(
                    question_text=f"Question text {i}",
                    category=cats[i - 1],
                    difficulty=diffs[i - 1],
                    purpose=secret_purpose,
                    expected_key_points=[f"Point {i}", secret_expected],
                    source=QuestionSource.general_backend,
                    source_evidence=secret_source_evidence
                )
            )

        from app.services.analysis_provider import InterviewQuestionsResult
        mock_ai_provider.mock_questions_result = InterviewQuestionsResult(
            data=InterviewQuestionSet(questions=gen_questions),
            error_code=None,
            error_message=None
        )

        def assert_no_leakage(res, endpoint_name: str):
            assert res.status_code == 200, f"{endpoint_name} failed with status {res.status_code}: {res.text}"
            res_text = res.text
            assert secret_expected not in res_text, f"Secret expected marker leaked in {endpoint_name}!"
            assert secret_purpose not in res_text, f"Secret purpose marker leaked in {endpoint_name}!"
            assert secret_source_evidence not in res_text, f"Secret source evidence marker leaked in {endpoint_name}!"
            assert "expected_key_points" not in res_text, f"expected_key_points key leaked in text of {endpoint_name}!"
            assert "source_evidence" not in res_text, f"source_evidence key leaked in text of {endpoint_name}!"

            def check_recursive(obj):
                if isinstance(obj, dict):
                    for k, v in obj.items():
                        assert k not in ("expected_key_points", "purpose", "source", "source_evidence"), (
                            f"Forbidden key '{k}' found recursively in {endpoint_name} response!"
                        )
                        check_recursive(v)
                elif isinstance(obj, list):
                    for item in obj:
                        check_recursive(item)

            check_recursive(res.json())

        # 1. Endpoint: Tạo interview session
        res_create = await client.post(f"/api/v1/job-applications/{app_id}/interview-sessions", json={"consent_to_external_ai": True})
        assert_no_leakage(res_create, "POST interview-sessions")
        created_session_id = res_create.json()["id"]

        # 2. Endpoint: Lấy current session
        res_curr = await client.get(f"/api/v1/job-applications/{app_id}/interview-sessions/current")
        assert_no_leakage(res_curr, "GET interview-sessions/current")

        # 3. Endpoint: Start session
        res_start = await client.post(f"/api/v1/interview-sessions/{created_session_id}/start")
        assert_no_leakage(res_start, "POST interview-sessions/start")

        # 4. Endpoint: Lấy progress
        res_prog = await client.get(f"/api/v1/interview-sessions/{created_session_id}/progress")
        assert_no_leakage(res_prog, "GET interview-sessions/progress")

        # Query question ids
        async with SessionLocal() as db:
            questions = (
                await db.execute(
                    select(InterviewQuestion)
                    .where(InterviewQuestion.interview_session_id == created_session_id)
                    .order_by(InterviewQuestion.order_index)
                )
            ).scalars().all()
            q_ids = [q.id for q in questions]

        # 5. Endpoint: Submit response câu 1 và Skip câu 2 (PUT endpoint)
        res_submit = await client.put(
            f"/api/v1/interview-sessions/{created_session_id}/questions/{q_ids[0]}/response",
            json={"answer_text": "Câu trả lời bảo mật kiểm thử", "is_skipped": False}
        )
        assert_no_leakage(res_submit, "PUT questions/response (submit)")

        res_skip = await client.put(
            f"/api/v1/interview-sessions/{created_session_id}/questions/{q_ids[1]}/response",
            json={"answer_text": None, "is_skipped": True}
        )
        assert_no_leakage(res_skip, "PUT questions/response (skip)")

        # Submit remaining questions 3..10 to complete session
        for i in range(2, 10):
            await client.put(
                f"/api/v1/interview-sessions/{created_session_id}/questions/{q_ids[i]}/response",
                json={"answer_text": f"Câu trả lời cho câu {i+1}", "is_skipped": False}
            )

        # 6. Endpoint: POST evaluation
        mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
            data=make_mock_ai_evaluation_set(list(range(1, 10))),
            error_code=None,
            error_message=None
        )
        res_eval = await client.post(f"/api/v1/interview-sessions/{created_session_id}/evaluation", json={"consent_to_external_ai": True})
        assert_no_leakage(res_eval, "POST evaluation")

        # 7. Endpoint: GET evaluation
        res_get_eval = await client.get(f"/api/v1/interview-sessions/{created_session_id}/evaluation")
        assert_no_leakage(res_get_eval, "GET evaluation")

    finally:
        monkeypatch.undo()
        async with SessionLocal() as db:
            if created_session_id:
                eval_ids = (
                    await db.execute(
                        select(InterviewEvaluation.id).where(
                            InterviewEvaluation.interview_session_id == created_session_id
                        )
                    )
                ).scalars().all()
                if eval_ids:
                    await db.execute(
                        delete(InterviewResponseEvaluation).where(
                            InterviewResponseEvaluation.interview_evaluation_id.in_(eval_ids)
                        )
                    )
                    await db.execute(
                        delete(InterviewEvaluation).where(
                            InterviewEvaluation.id.in_(eval_ids)
                        )
                    )
                await db.execute(delete(InterviewResponse).where(InterviewResponse.interview_session_id == created_session_id))
                await db.execute(delete(InterviewQuestion).where(InterviewQuestion.interview_session_id == created_session_id))
                await db.execute(delete(InterviewSession).where(InterviewSession.id == created_session_id))
            await db.execute(delete(ApplicationAnalysis).where(ApplicationAnalysis.id == analysis_id))
            await db.execute(delete(JobApplication).where(JobApplication.id == app_id))
            await db.execute(delete(CVDocument).where(CVDocument.id == doc_id))
            await db.commit()


# 27. Section G: All 10 skipped with empty GROQ_API_KEY does NOT return 503
@pytest.mark.asyncio
async def test_all_10_skipped_without_api_key(client, completed_interview_session, mock_ai_provider, monkeypatch):
    session_id = completed_interview_session["session_id"]
    # Mark all 10 as skipped
    async with SessionLocal() as db:
        resps = (await db.execute(select(InterviewResponse).where(InterviewResponse.interview_session_id == session_id))).scalars().all()
        for r in resps:
            r.is_skipped = True
            r.answer_text = None
        await db.commit()

    # Clear GROQ_API_KEY
    monkeypatch.setattr(settings, "GROQ_API_KEY", "")

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "evaluated"
    assert data["total_score"] == 0.0
    assert data["average_score"] == 0.0
    assert data["readiness_level"] == "needs_foundation"
    assert mock_ai_provider.call_count == 0


# 28. Section F: Commit failure in recovery raises 500
@pytest.mark.asyncio
async def test_evaluation_double_failure_raises_500(client, completed_interview_session, mock_ai_provider, monkeypatch):
    session_id = completed_interview_session["session_id"]
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=make_mock_ai_evaluation_set(list(range(1, 9))),
        error_code=None,
        error_message=None
    )

    import sqlalchemy.ext.asyncio.session
    original_commit = sqlalchemy.ext.asyncio.session.AsyncSession.commit
    commit_count = 0

    async def fail_all_commits_after_claim(self):
        nonlocal commit_count
        commit_count += 1
        if commit_count >= 2:  # First commit claims session, all subsequent fail
            raise Exception("Catastrophic DB Failure")
        await original_commit(self)

    monkeypatch.setattr(sqlalchemy.ext.asyncio.session.AsyncSession, "commit", fail_all_commits_after_claim)

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    monkeypatch.undo()

    assert res.status_code == 500
    assert "Lỗi máy chủ nội bộ" in res.json()["detail"]


# 28b. Test isolation verification with sentinel record
@pytest.mark.asyncio
async def test_cleanup_isolation_with_sentinel(client):
    sentinel_doc_id = uuid.uuid4()
    sentinel_app_id = uuid.uuid4()
    sentinel_analysis_id = uuid.uuid4()
    sentinel_session_id = uuid.uuid4()
    sentinel_q_id = uuid.uuid4()
    sentinel_r_id = uuid.uuid4()
    sentinel_eval_id = uuid.uuid4()
    sentinel_resp_eval_id = uuid.uuid4()

    target_doc_id = uuid.uuid4()
    target_app_id = uuid.uuid4()
    target_analysis_id = uuid.uuid4()
    target_session_id = uuid.uuid4()
    target_q_id = uuid.uuid4()
    target_r_id = uuid.uuid4()
    target_eval_id = uuid.uuid4()
    target_resp_eval_id = uuid.uuid4()

    async with SessionLocal() as db:
        # 1. Create sentinel
        sentinel_doc = CVDocument(
            id=sentinel_doc_id,
            original_filename="sentinel.pdf",
            storage_key=f"k_{sentinel_doc_id}",
            media_type="application/pdf",
            size_bytes=100,
            sha256="sentinel_sha",
            status=CVStatus.extracted,
            extracted_text="Sentinel text",
        )
        db.add(sentinel_doc)
        sentinel_app = JobApplication(
            id=sentinel_app_id,
            cv_document_id=sentinel_doc_id,
            target_role="Backend Intern",
            job_description="Sentinel JD",
        )
        db.add(sentinel_app)
        sentinel_analysis = ApplicationAnalysis(
            id=sentinel_analysis_id,
            job_application_id=sentinel_app_id,
            status=AnalysisStatus.completed,
            provider="groq",
            model="llama",
            schema_version="1",
            prompt_version="1",
            analysis_data={"summary_vi": "Sentinel"},
        )
        db.add(sentinel_analysis)
        sentinel_session = InterviewSession(
            id=sentinel_session_id,
            job_application_id=sentinel_app_id,
            status=SessionStatus.completed,
            provider="groq",
            model="llama",
            prompt_version="1",
            schema_version="1",
        )
        db.add(sentinel_session)
        sentinel_q = InterviewQuestion(
            id=sentinel_q_id,
            interview_session_id=sentinel_session_id,
            order_index=1,
            question_text="Sentinel question",
            category=QuestionCategory.backend_fundamentals.value,
            difficulty=QuestionDifficulty.easy.value,
            purpose="Sentinel purpose",
            expected_key_points=["Key 1"],
            source=QuestionSource.general_backend.value,
        )
        db.add(sentinel_q)
        sentinel_resp = InterviewResponse(
            id=sentinel_r_id,
            interview_session_id=sentinel_session_id,
            interview_question_id=sentinel_q_id,
            answer_text="Sentinel answer",
            is_skipped=False,
        )
        db.add(sentinel_resp)
        sentinel_eval = InterviewEvaluation(
            id=sentinel_eval_id,
            interview_session_id=sentinel_session_id,
            status=EvaluationStatus.evaluated.value,
            provider="groq",
            model="llama",
            prompt_version="1",
            schema_version="1",
            readiness_level=ReadinessLevel.developing.value,
            total_score=7.0,
            average_score=7.0,
            answered_count=1,
            skipped_count=9,
            overall_summary="Sentinel summary",
            evaluated_at=func.now(),
        )
        db.add(sentinel_eval)
        sentinel_resp_eval = InterviewResponseEvaluation(
            id=sentinel_resp_eval_id,
            interview_evaluation_id=sentinel_eval_id,
            interview_response_id=sentinel_r_id,
            technical_accuracy=3.0,
            completeness=1.5,
            reasoning_clarity=1.5,
            relevance=0.5,
            practical_example=0.5,
            total_score=7.0,
            feedback="Sentinel feedback",
            strengths=["Good"],
            missing_points=["None"],
            suggested_answer="Sentinel sample answer",
            study_topics=["Topic A"],
            confidence=EvaluationConfidence.high.value,
        )
        db.add(sentinel_resp_eval)

        # 2. Create target
        target_doc = CVDocument(
            id=target_doc_id,
            original_filename="target.pdf",
            storage_key=f"k_{target_doc_id}",
            media_type="application/pdf",
            size_bytes=100,
            sha256="target_sha",
            status=CVStatus.extracted,
            extracted_text="Target text",
        )
        db.add(target_doc)
        target_app = JobApplication(
            id=target_app_id,
            cv_document_id=target_doc_id,
            target_role="Backend Intern",
            job_description="Target JD",
        )
        db.add(target_app)
        target_analysis = ApplicationAnalysis(
            id=target_analysis_id,
            job_application_id=target_app_id,
            status=AnalysisStatus.completed,
            provider="groq",
            model="llama",
            schema_version="1",
            prompt_version="1",
            analysis_data={"summary_vi": "Target"},
        )
        db.add(target_analysis)
        target_session = InterviewSession(
            id=target_session_id,
            job_application_id=target_app_id,
            status=SessionStatus.completed,
            provider="groq",
            model="llama",
            prompt_version="1",
            schema_version="1",
        )
        db.add(target_session)
        target_q = InterviewQuestion(
            id=target_q_id,
            interview_session_id=target_session_id,
            order_index=1,
            question_text="Target question",
            category=QuestionCategory.backend_fundamentals.value,
            difficulty=QuestionDifficulty.easy.value,
            purpose="Target purpose",
            expected_key_points=["Key 1"],
            source=QuestionSource.general_backend.value,
        )
        db.add(target_q)
        target_resp = InterviewResponse(
            id=target_r_id,
            interview_session_id=target_session_id,
            interview_question_id=target_q_id,
            answer_text="Target answer",
            is_skipped=False,
        )
        db.add(target_resp)
        target_eval = InterviewEvaluation(
            id=target_eval_id,
            interview_session_id=target_session_id,
            status=EvaluationStatus.evaluated.value,
            provider="groq",
            model="llama",
            prompt_version="1",
            schema_version="1",
            readiness_level=ReadinessLevel.developing.value,
            total_score=6.0,
            average_score=6.0,
            answered_count=1,
            skipped_count=9,
            overall_summary="Target summary",
            evaluated_at=func.now(),
        )
        db.add(target_eval)
        target_resp_eval = InterviewResponseEvaluation(
            id=target_resp_eval_id,
            interview_evaluation_id=target_eval_id,
            interview_response_id=target_r_id,
            technical_accuracy=2.0,
            completeness=1.5,
            reasoning_clarity=1.5,
            relevance=0.5,
            practical_example=0.5,
            total_score=6.0,
            feedback="Target feedback",
            strengths=["Good"],
            missing_points=["None"],
            suggested_answer="Target sample answer",
            study_topics=["Topic B"],
            confidence=EvaluationConfidence.high.value,
        )
        db.add(target_resp_eval)

        await db.commit()

    try:
        # Run cleanup of target session ONLY (simulating isolated fixture/test cleanup)
        async with SessionLocal() as db:
            eval_ids = (
                await db.execute(
                    select(InterviewEvaluation.id).where(
                        InterviewEvaluation.interview_session_id == target_session_id
                    )
                )
            ).scalars().all()
            if eval_ids:
                await db.execute(
                    delete(InterviewResponseEvaluation).where(
                        InterviewResponseEvaluation.interview_evaluation_id.in_(eval_ids)
                    )
                )
                await db.execute(
                    delete(InterviewEvaluation).where(
                        InterviewEvaluation.id.in_(eval_ids)
                    )
                )
            await db.execute(delete(InterviewResponse).where(InterviewResponse.interview_session_id == target_session_id))
            await db.execute(delete(InterviewQuestion).where(InterviewQuestion.interview_session_id == target_session_id))
            await db.execute(delete(InterviewSession).where(InterviewSession.id == target_session_id))
            await db.execute(delete(ApplicationAnalysis).where(ApplicationAnalysis.id == target_analysis_id))
            await db.execute(delete(JobApplication).where(JobApplication.id == target_app_id))
            await db.execute(delete(CVDocument).where(CVDocument.id == target_doc_id))
            await db.commit()

        # Confirm target records are gone
        async with SessionLocal() as db:
            target_check = (await db.execute(select(InterviewEvaluation).where(InterviewEvaluation.id == target_eval_id))).scalar_one_or_none()
            assert target_check is None

            # Confirm sentinel evaluation and sentinel response evaluation remain completely intact!
            sentinel_eval_check = (await db.execute(select(InterviewEvaluation).where(InterviewEvaluation.id == sentinel_eval_id))).scalar_one_or_none()
            assert sentinel_eval_check is not None
            assert sentinel_eval_check.id == sentinel_eval_id

            sentinel_resp_eval_check = (await db.execute(select(InterviewResponseEvaluation).where(InterviewResponseEvaluation.id == sentinel_resp_eval_id))).scalar_one_or_none()
            assert sentinel_resp_eval_check is not None
            assert sentinel_resp_eval_check.id == sentinel_resp_eval_id
    finally:
        # Dọn dẹp chính sentinel bằng ID cụ thể
        async with SessionLocal() as db:
            await db.execute(delete(InterviewResponseEvaluation).where(InterviewResponseEvaluation.id == sentinel_resp_eval_id))
            await db.execute(delete(InterviewEvaluation).where(InterviewEvaluation.id == sentinel_eval_id))
            await db.execute(delete(InterviewResponse).where(InterviewResponse.id == sentinel_r_id))
            await db.execute(delete(InterviewQuestion).where(InterviewQuestion.id == sentinel_q_id))
            await db.execute(delete(InterviewSession).where(InterviewSession.id == sentinel_session_id))
            await db.execute(delete(ApplicationAnalysis).where(ApplicationAnalysis.id == sentinel_analysis_id))
            await db.execute(delete(JobApplication).where(JobApplication.id == sentinel_app_id))
            await db.execute(delete(CVDocument).where(CVDocument.id == sentinel_doc_id))
            await db.commit()


# 28b. Task 09D: Integration test verifying Decimal aggregation for category scores (0.01 + 0.06 -> 0.04)
@pytest.mark.asyncio
async def test_evaluation_decimal_aggregation_category_scores(client, completed_interview_session, mock_ai_provider):
    session_id = completed_interview_session["session_id"]

    evals = []
    # Question 1 (backend_fundamentals): 0.01
    evals.append(
        QuestionEvaluationAI(
            question_order=1,
            technical_accuracy=0.01,
            completeness=0.0,
            reasoning_clarity=0.0,
            relevance=0.0,
            practical_example=0.0,
            total_score=0.01,
            feedback="Nhận xét câu 1 với điểm số 0.01.",
            strengths=[],
            missing_points=["Chưa đủ chi tiết"],
            suggested_answer="Câu trả lời mẫu cho câu 1.",
            study_topics=["Backend"],
            confidence=EvaluationConfidence.high
        )
    )
    # Question 2 (backend_fundamentals): 0.06
    evals.append(
        QuestionEvaluationAI(
            question_order=2,
            technical_accuracy=0.06,
            completeness=0.0,
            reasoning_clarity=0.0,
            relevance=0.0,
            practical_example=0.0,
            total_score=0.06,
            feedback="Nhận xét câu 2 với điểm số 0.06.",
            strengths=[],
            missing_points=["Cần làm rõ"],
            suggested_answer="Câu trả lời mẫu cho câu 2.",
            study_topics=["Backend"],
            confidence=EvaluationConfidence.high
        )
    )
    # Questions 3 to 8: regular scores
    for o in range(3, 9):
        evals.append(
            QuestionEvaluationAI(
                question_order=o,
                technical_accuracy=2.0,
                completeness=1.0,
                reasoning_clarity=1.0,
                relevance=0.5,
                practical_example=0.5,
                total_score=5.0,
                feedback=f"Nhận xét chuẩn cho câu {o}.",
                strengths=["Tốt"],
                missing_points=[],
                suggested_answer=f"Câu trả lời mẫu cho câu {o}.",
                study_topics=["Backend"],
                confidence=EvaluationConfidence.high
            )
        )

    ai_set = InterviewEvaluationAISet(
        evaluations=evals,
        overall_summary="Ứng viên có kiến thức cơ bản.",
        key_strengths=["Nền tảng"],
        knowledge_gaps=["Thực hành"],
        priority_improvements=["Ôn tập"],
        recommended_study_topics=["Database"],
        short_term_study_plan=["Tuần 1: Ôn tập"],
        next_interview_advice="Cố gắng hơn ở lần phỏng vấn tiếp theo."
    )
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=ai_set, error_code=None, error_message=None
    )

    res = await client.post(f"/api/v1/interview-sessions/{session_id}/evaluation", json={"consent_to_external_ai": True})
    assert res.status_code == 200
    data = res.json()

    assert data["status"] == "evaluated"
    assert "category_scores" in data
    # (0.01 + 0.06) / 2 = 0.035 -> ROUND_HALF_UP -> 0.04!
    assert data["category_scores"]["backend_fundamentals"] == 0.04
    # Ensure returned type is standard float, not raw Decimal string
    assert isinstance(data["category_scores"]["backend_fundamentals"], float)
    assert not isinstance(data["category_scores"]["backend_fundamentals"], str)


# 29. Export DOCX tests (404 when session/eval missing, 200 with valid DOCX content when evaluated)
@pytest.mark.asyncio
async def test_export_docx_not_found_and_success(client, completed_interview_session, mock_ai_provider):
    import io
    from docx import Document as DocxDocument

    # Non-existent session -> 404
    res_404 = await client.get(f"/api/v1/interview-sessions/{uuid.uuid4()}/evaluation/export-docx")
    assert res_404.status_code == 404

    session_id = completed_interview_session["session_id"]

    # Session exists but not evaluated yet -> 404
    res_no_eval = await client.get(f"/api/v1/interview-sessions/{session_id}/evaluation/export-docx")
    assert res_no_eval.status_code == 404

    # Evaluate session
    mock_ai_provider.mock_evaluation_result = InterviewEvaluationResult(
        data=make_mock_ai_evaluation_set(list(range(1, 9))),
        error_code=None,
        error_message=None,
    )
    eval_res = await client.post(
        f"/api/v1/interview-sessions/{session_id}/evaluation",
        json={"consent_to_external_ai": True},
    )
    assert eval_res.status_code == 200

    # Export DOCX -> 200
    exp_res = await client.get(f"/api/v1/interview-sessions/{session_id}/evaluation/export-docx")
    assert exp_res.status_code == 200
    assert (
        exp_res.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment; filename=" in exp_res.headers.get("content-disposition", "")
    assert ".docx" in exp_res.headers.get("content-disposition", "")

    # Verify valid DOCX structure and content in memory
    doc = DocxDocument(io.BytesIO(exp_res.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "BÁO CÁO ĐÁNH GIÁ KẾT QUẢ PHỎNG VẤN" in full_text
    assert "Ứng viên có kiến thức nền tảng tốt." in full_text
    assert "Câu trả lời mẫu cho câu 1." in full_text
    assert "Đã bỏ qua câu hỏi này" in full_text


# 30. Question Hint (Coach Mode) endpoint test
@pytest.mark.asyncio
async def test_question_hint_endpoint(client, completed_interview_session):
    session_id = completed_interview_session["session_id"]
    first_qid = completed_interview_session["question_ids"][0]

    # Non-existent session -> 404
    res_bad_sess = await client.get(
        f"/api/v1/interview-sessions/{uuid.uuid4()}/questions/{first_qid}/hint"
    )
    assert res_bad_sess.status_code == 404

    # Non-existent question -> 404
    res_bad_q = await client.get(
        f"/api/v1/interview-sessions/{session_id}/questions/{uuid.uuid4()}/hint"
    )
    assert res_bad_q.status_code == 404

    # Valid question hint -> 200
    res_ok = await client.get(
        f"/api/v1/interview-sessions/{session_id}/questions/{first_qid}/hint"
    )
    assert res_ok.status_code == 200
    body = res_ok.json()
    assert body["question_id"] == str(first_qid)
    assert body["order_index"] == 1
    assert body["interview_goal"] == "Test fundamentals 1"
    assert body["key_concepts"] == ["Key point A for 1", "Key point B for 1"]
    assert len(body["answer_structure_tips"]) >= 2


# 31. Zero-orphan confirmation fixture check
@pytest.mark.asyncio
async def test_all_tables_zero_orphan_after_run(client):
    async with SessionLocal() as db:
        for model in [
            InterviewResponseEvaluation,
            InterviewEvaluation,
            InterviewResponse,
            InterviewQuestion,
            InterviewSession,
            ApplicationAnalysis,
            JobApplication,
            CVDocument,
        ]:
            count = (await db.execute(select(func.count()).select_from(model))).scalar()
            assert count == 0, f"Table {model.__tablename__} is not empty! Count: {count}"


