import pytest
import io
from pathlib import Path
from pypdf import PdfWriter, PdfReader
import docx
from app.services.extractor import PDFExtractor, DOCXExtractor, CVTextExtractor, ExtractionError, MAX_TEXT_LENGTH
import tempfile
import asyncio

def create_pdf(text: str = "", pages: int = 1, encrypt_pw: str = None) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=100, height=100)
        
    if encrypt_pw:
        writer.encrypt(encrypt_pw)
        
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()

def create_docx(text_elements: list) -> bytes:
    doc = docx.Document()
    for el in text_elements:
        if isinstance(el, str):
            doc.add_paragraph(el)
        elif isinstance(el, list):
            # table
            table = doc.add_table(rows=1, cols=len(el))
            row_cells = table.rows[0].cells
            for i, val in enumerate(el):
                row_cells[i].text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

@pytest.fixture
def tmp_pdf_path(tmp_path):
    pdf = create_pdf(pages=1)
    p = tmp_path / "test.pdf"
    p.write_bytes(pdf)
    return p

def test_pdf_extract_sync_empty(tmp_pdf_path):
    with pytest.raises(ExtractionError, match="Không tìm thấy nội dung văn bản"):
        PDFExtractor.extract_sync(tmp_pdf_path)

def test_pdf_extract_sync_with_text(tmp_pdf_path, monkeypatch):
    def mock_extract_text(self):
        return "Xin chào Việt Nam \x00\r\n Test"
    monkeypatch.setattr("pypdf.PageObject.extract_text", mock_extract_text)
    
    res = PDFExtractor.extract_sync(tmp_pdf_path)
    assert res == "Xin chào Việt Nam \n Test"

def test_pdf_encrypted(tmp_path):
    pdf = create_pdf(pages=1, encrypt_pw="secret")
    p = tmp_path / "enc.pdf"
    p.write_bytes(pdf)
    with pytest.raises(ExtractionError, match="Không thể đọc PDF bị khóa mật khẩu"):
        PDFExtractor.extract_sync(p)

def test_pdf_too_many_pages(tmp_path):
    pdf = create_pdf(pages=51)
    p = tmp_path / "big.pdf"
    p.write_bytes(pdf)
    with pytest.raises(ExtractionError, match="PDF vượt quá giới hạn 50 trang"):
        PDFExtractor.extract_sync(p)

def test_docx_extract_sync_empty(tmp_path):
    docx_bytes = create_docx([])
    p = tmp_path / "empty.docx"
    p.write_bytes(docx_bytes)
    with pytest.raises(ExtractionError, match="Không tìm thấy nội dung văn bản trong DOCX"):
        DOCXExtractor.extract_sync(p)

def test_docx_extract_sync_with_text(tmp_path):
    docx_bytes = create_docx(["Đoạn 1", ["Cột 1", "Cột 2"], "Đoạn 2"])
    p = tmp_path / "test.docx"
    p.write_bytes(docx_bytes)
    res = DOCXExtractor.extract_sync(p)
    assert "Đoạn 1\nCột 1 | Cột 2\nĐoạn 2" in res

def test_text_normalization():
    raw = "Hello\r\nWorld\r\r\n\n\n\nSpace   \x00\x08Tab\t!"
    res = PDFExtractor._normalize_text(raw)
    assert res == "Hello\nWorld\n\nSpace Tab !"

def test_text_length_limit():
    raw = "a" * (MAX_TEXT_LENGTH + 1)
    with pytest.raises(ExtractionError, match="vượt quá giới hạn"):
        PDFExtractor._validate_length(raw)

@pytest.mark.asyncio
async def test_cv_text_extractor_unsupported():
    with pytest.raises(ExtractionError, match="không được hỗ trợ"):
        await CVTextExtractor.extract_text(Path("test.txt"), "text/plain")
