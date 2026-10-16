import io
import uuid
from datetime import timezone
from typing import Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.interview_evaluation import (
    EvaluationStatus,
    InterviewEvaluation,
    InterviewResponseEvaluation,
    ReadinessLevel,
)
from app.models.interview_response import InterviewResponse
from app.models.interview_session import InterviewSession
from app.models.job_application import JobApplication


READINESS_LABELS = {
    "strong_intern_candidate": "Ứng viên Intern xuất sắc (Strong Intern Candidate)",
    "intern_ready": "Sẵn sàng làm Intern (Intern Ready)",
    "developing": "Đang phát triển (Developing)",
    "needs_foundation": "Cần củng cố nền tảng (Needs Foundation)",
}

CATEGORY_LABELS = {
    "backend_fundamentals": "Kiến thức nền tảng Backend",
    "cv_project_experience": "Kinh nghiệm dự án & CV",
    "database": "Cơ sở dữ liệu (Database)",
    "api_web_design": "Thiết kế API & Web",
    "debugging_problem_solving": "Gỡ lỗi & Giải quyết vấn đề",
    "behavioral_teamwork": "Kỹ năng mềm & Làm việc nhóm",
}

DIFFICULTY_LABELS = {
    "easy": "Cơ bản (Easy)",
    "medium": "Trung bình (Medium)",
    "hard": "Nâng cao (Hard)",
}


def _format_score(val) -> str:
    if val is None:
        return "0.0"
    return f"{float(val):.2f}".rstrip("0").rstrip(".") if "." in f"{float(val):.2f}" else str(val)


class EvaluationReportExporter:
    @staticmethod
    async def export_docx(
        db: AsyncSession, session_id: uuid.UUID
    ) -> Tuple[bytes, str]:
        session = (
            await db.execute(
                select(InterviewSession).where(InterviewSession.id == session_id)
            )
        ).scalars().first()
        if not session:
            raise HTTPException(
                status_code=404, detail="Không tìm thấy phiên phỏng vấn."
            )

        eval_record = (
            await db.execute(
                select(InterviewEvaluation).where(
                    InterviewEvaluation.interview_session_id == session_id
                )
            )
        ).scalars().first()
        if not eval_record:
            raise HTTPException(
                status_code=404,
                detail="Không tìm thấy báo cáo đánh giá cho buổi phỏng vấn này.",
            )

        if eval_record.status != EvaluationStatus.evaluated:
            raise HTTPException(
                status_code=409,
                detail="Báo cáo đánh giá chưa hoàn tất, không thể xuất file.",
            )

        job_app = (
            await db.execute(
                select(JobApplication).where(
                    JobApplication.id == session.job_application_id
                )
            )
        ).scalars().first()

        resp_evals = (
            await db.execute(
                select(InterviewResponseEvaluation)
                .options(
                    selectinload(InterviewResponseEvaluation.response).selectinload(
                        InterviewResponse.question
                    )
                )
                .where(
                    InterviewResponseEvaluation.interview_evaluation_id
                    == eval_record.id
                )
            )
        ).scalars().all()

        resp_evals = sorted(
            resp_evals, key=lambda x: x.response.question.order_index
        )

        doc = Document()

        # Configure default normal style
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Arial"
        font.size = Pt(11)

        # Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run("BÁO CÁO ĐÁNH GIÁ KẾT QUẢ PHỎNG VẤN")
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.color.rgb = RGBColor(29, 78, 216)

        subtitle_p = doc.add_paragraph()
        subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        target_role = job_app.target_role if job_app else "Backend Developer Intern"
        evaluated_str = (
            eval_record.evaluated_at.astimezone(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
            if eval_record.evaluated_at
            else "N/A"
        )
        sub_run = subtitle_p.add_run(
            f"Vị trí: {target_role}  |  Lần luyện số: {session.attempt_number}  |  Ngày đánh giá: {evaluated_str}"
        )
        sub_run.font.size = Pt(10)
        sub_run.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph()

        # 1. TỔNG QUAN KẾT QUẢ
        h1 = doc.add_heading("1. Tổng quan kết quả", level=1)
        for run in h1.runs:
            run.font.color.rgb = RGBColor(15, 23, 42)

        summary_table = doc.add_table(rows=2, cols=4)
        summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        summary_table.style = "Table Grid"

        headers = ["Tổng điểm", "Điểm trung bình", "Số câu đã trả lời", "Số câu bỏ qua"]
        values = [
            f"{_format_score(eval_record.total_score)} / 100",
            f"{_format_score(eval_record.average_score)} / 10",
            f"{eval_record.answered_count or 0} / 10",
            f"{eval_record.skipped_count or 0} / 10",
        ]

        for col_idx in range(4):
            cell_h = summary_table.cell(0, col_idx)
            cell_h.text = headers[col_idx]
            for p in cell_h.paragraphs:
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(10)

            cell_v = summary_table.cell(1, col_idx)
            cell_v.text = values[col_idx]
            for p in cell_v.paragraphs:
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(12)
                    r.font.color.rgb = RGBColor(29, 78, 216)

        readiness_p = doc.add_paragraph()
        readiness_p.paragraph_format.space_before = Pt(10)
        r_label = readiness_p.add_run("Mức độ sẵn sàng: ")
        r_label.bold = True
        rl_key = (
            getattr(eval_record.readiness_level, "value", str(eval_record.readiness_level))
            if eval_record.readiness_level
            else "N/A"
        )
        r_val = readiness_p.add_run(READINESS_LABELS.get(rl_key, rl_key))
        r_val.bold = True
        r_val.font.color.rgb = RGBColor(5, 150, 105)

        # 2. ĐIỂM THEO NHÓM CHỦ ĐỀ
        if eval_record.category_scores:
            h2 = doc.add_heading("2. Điểm trung bình theo nhóm chủ đề", level=1)
            for run in h2.runs:
                run.font.color.rgb = RGBColor(15, 23, 42)

            cat_items = list(eval_record.category_scores.items())
            cat_table = doc.add_table(rows=len(cat_items) + 1, cols=2)
            cat_table.style = "Table Grid"
            cat_table.cell(0, 0).text = "Nhóm chủ đề"
            cat_table.cell(0, 1).text = "Điểm trung bình (thang 10)"
            for col_idx in range(2):
                for p in cat_table.cell(0, col_idx).paragraphs:
                    for r in p.runs:
                        r.bold = True

            for idx, (cat_key, cat_score) in enumerate(cat_items, start=1):
                cat_table.cell(idx, 0).text = CATEGORY_LABELS.get(
                    cat_key, cat_key.replace("_", " ").title()
                )
                cat_table.cell(idx, 1).text = f"{_format_score(cat_score)} / 10"

        # 3. NHẬN XÉT TỔNG THỂ & LỘ TRÌNH HỌC TẬP
        h3 = doc.add_heading("3. Nhận xét tổng thể & Lộ trình cải thiện", level=1)
        for run in h3.runs:
            run.font.color.rgb = RGBColor(15, 23, 42)

        if eval_record.overall_summary:
            p_sum_title = doc.add_paragraph()
            r_sum = p_sum_title.add_run("Nhận xét tổng quát:")
            r_sum.bold = True
            doc.add_paragraph(eval_record.overall_summary)

        def _add_bullet_section(title: str, items: list | None):
            if not items:
                return
            p_t = doc.add_paragraph()
            p_t.paragraph_format.space_before = Pt(6)
            run_t = p_t.add_run(title)
            run_t.bold = True
            for item in items:
                doc.add_paragraph(str(item), style="List Bullet")

        _add_bullet_section("Điểm mạnh nổi bật:", eval_record.key_strengths)
        _add_bullet_section("Lỗ hổng kiến thức cần lưu ý:", eval_record.knowledge_gaps)
        _add_bullet_section("Ưu tiên cải thiện hàng đầu:", eval_record.priority_improvements)
        _add_bullet_section("Chủ đề nên ôn tập:", eval_record.recommended_study_topics)
        _add_bullet_section("Kế hoạch học tập ngắn hạn:", eval_record.short_term_study_plan)

        if eval_record.next_interview_advice:
            p_adv_title = doc.add_paragraph()
            p_adv_title.paragraph_format.space_before = Pt(6)
            r_adv = p_adv_title.add_run("Lời khuyên cho buổi phỏng vấn tiếp theo:")
            r_adv.bold = True
            p_adv = doc.add_paragraph(eval_record.next_interview_advice)
            for r in p_adv.runs:
                r.italic = True

        # 4. CHI TIẾT TỪNG CÂU HỎI (1 - 10)
        h4 = doc.add_heading("4. Chi tiết đánh giá từng câu hỏi (10 câu)", level=1)
        for run in h4.runs:
            run.font.color.rgb = RGBColor(15, 23, 42)

        for re_item in resp_evals:
            q = re_item.response.question
            r = re_item.response
            cat_str = CATEGORY_LABELS.get(
                q.category.value if hasattr(q.category, "value") else str(q.category),
                str(q.category),
            )
            diff_str = DIFFICULTY_LABELS.get(
                q.difficulty.value if hasattr(q.difficulty, "value") else str(q.difficulty),
                str(q.difficulty),
            )

            qh = doc.add_heading(
                f"Câu {q.order_index}: {_format_score(re_item.total_score)}/10 điểm — [{cat_str} | {diff_str}]",
                level=2,
            )
            for run in qh.runs:
                run.font.color.rgb = RGBColor(29, 78, 216)

            # Question text
            pq = doc.add_paragraph()
            rq_label = pq.add_run("Câu hỏi: ")
            rq_label.bold = True
            pq.add_run(q.question_text)

            # Candidate answer
            pa = doc.add_paragraph()
            ra_label = pa.add_run("Câu trả lời của ứng viên: ")
            ra_label.bold = True
            if r.is_skipped:
                ra_val = pa.add_run("(Đã bỏ qua câu hỏi này)")
                ra_val.italic = True
            else:
                pa.add_run(r.answer_text or "")

            # Rubric breakdown line
            pr = doc.add_paragraph()
            rr_label = pr.add_run("Điểm thành phần Rubric: ")
            rr_label.bold = True
            pr.add_run(
                f"Kỹ thuật: {_format_score(re_item.technical_accuracy)}/4.0  |  "
                f"Đầy đủ: {_format_score(re_item.completeness)}/2.0  |  "
                f"Lập luận: {_format_score(re_item.reasoning_clarity)}/2.0  |  "
                f"Trọng tâm: {_format_score(re_item.relevance)}/1.0  |  "
                f"Ví dụ thực tế: {_format_score(re_item.practical_example)}/1.0"
            )

            # Feedback
            pf = doc.add_paragraph()
            rf_label = pf.add_run("Nhận xét của AI: ")
            rf_label.bold = True
            pf.add_run(re_item.feedback or "")

            # Strengths & Missing points
            if re_item.strengths:
                _add_bullet_section("Điểm tốt:", re_item.strengths)
            if re_item.missing_points:
                _add_bullet_section("Điểm còn thiếu:", re_item.missing_points)

            # Suggested answer
            if re_item.suggested_answer:
                ps = doc.add_paragraph()
                rs_label = ps.add_run("Câu trả lời mẫu gợi ý: ")
                rs_label.bold = True
                ps.add_run(re_item.suggested_answer)

            # Study topics
            if re_item.study_topics:
                pt = doc.add_paragraph()
                rt_label = pt.add_run("Chủ đề cần ôn tập: ")
                rt_label.bold = True
                pt.add_run(", ".join(re_item.study_topics))

        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()

        short_id = str(session_id)[:8]
        filename = f"Bao_Cao_Phong_Van_Lan_{session.attempt_number}_{short_id}.docx"
        return docx_bytes, filename
