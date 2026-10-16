from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def paragraph_payload(paragraph, index: int) -> dict:
    ppr = paragraph._p.pPr
    num_id = None
    ilvl = None
    if ppr is not None and ppr.numPr is not None:
        if ppr.numPr.numId is not None:
            num_id = ppr.numPr.numId.val
        if ppr.numPr.ilvl is not None:
            ilvl = ppr.numPr.ilvl.val
    return {
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text,
        "alignment": str(paragraph.alignment),
        "num_id": num_id,
        "ilvl": ilvl,
        "runs": [
            {
                "text": run.text,
                "bold": run.bold,
                "italic": run.italic,
                "underline": run.underline,
                "size_pt": run.font.size.pt if run.font.size else None,
                "font": run.font.name,
            }
            for run in paragraph.runs
        ],
    }


def cell_payload(cell) -> dict:
    return {
        "text": cell.text,
        "paragraphs": [paragraph_payload(p, i) for i, p in enumerate(cell.paragraphs)],
    }


def main() -> None:
    path = Path(sys.argv[1])
    summary_only = "--summary" in sys.argv[2:]
    doc = Document(path)
    if summary_only:
        for i, paragraph in enumerate(doc.paragraphs):
            print(f"P{i:02d} [{paragraph.style.name if paragraph.style else ''}] {paragraph.text}")
        for ti, table in enumerate(doc.tables):
            print(f"TABLE {ti}: {len(table.rows)} rows x {len(table.columns)} cols")
            for ri, row in enumerate(table.rows):
                cells = [cell.text.replace(chr(10), " / ") for cell in row.cells]
                print(f"  R{ri:02d}: " + " || ".join(cells))
        return
    payload = {
        "paragraphs": [paragraph_payload(p, i) for i, p in enumerate(doc.paragraphs)],
        "tables": [
            {
                "index": ti,
                "rows": [
                    [cell_payload(cell) for cell in row.cells]
                    for row in table.rows
                ],
            }
            for ti, table in enumerate(doc.tables)
        ],
        "sections": [],
        "styles": [],
        "inline_shapes": len(doc.inline_shapes),
        "all_text_nodes": [node.text for node in doc.element.body.iter(qn("w:t"))],
    }
    for i, section in enumerate(doc.sections):
        payload["sections"].append(
            {
                "index": i,
                "page_width": section.page_width,
                "page_height": section.page_height,
                "top_margin": section.top_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
                "right_margin": section.right_margin,
                "header_distance": section.header_distance,
                "footer_distance": section.footer_distance,
                "header_text": "\n".join(p.text for p in section.header.paragraphs),
                "footer_text": "\n".join(p.text for p in section.footer.paragraphs),
            }
        )
    for style in doc.styles:
        if style.type is None:
            continue
        payload["styles"].append(
            {
                "name": style.name,
                "type": str(style.type),
                "builtin": style.builtin,
            }
        )
    json.dump(payload, sys.stdout, ensure_ascii=False, indent=2, default=str)


if __name__ == "__main__":
    main()
